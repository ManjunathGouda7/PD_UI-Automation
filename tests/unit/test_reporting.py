"""Unit tests for Scripts.Helpers.reporting (naming, history, report files)."""
import json
from pathlib import Path

import openpyxl
import pytest

from Scripts.Helpers import reporting

TEST_DIR = Path(__file__).resolve().parents[2] / "Scripts" / "TestCases"


@pytest.mark.parametrize("name, expected", [
    ("test_TC_01_tester_info_table_has_all_fields", "TC_01 · Tester info table has all fields"),
    ("test_TC_02_ip_field_matches_tester_ip_row_and_history", "TC_02 · IP field matches tester IP row and history"),
    ("test_TC_03_build_tip", "TC_03 · Build tip"),  # 'tip' must not become 'tIP'
    ("test_server_online", "Server online"),
    ("test_TC_04_x[#chart_toolbar_fit_button]", "TC_04 · X [#chart_toolbar_fit_button]"),
])
def test_humanize_test_name(name, expected):
    assert reporting.humanize_test_name(name) == expected


@pytest.mark.parametrize("title, expected", [
    ("Panel 1: Connection Setup", "Connection Setup"),
    ("REST API: read-only backend checks", "REST API"),
    ("Application shell: load, header", "Application shell"),
    ("No colon here", "No colon here"),
])
def test_short_name_from_title(title, expected):
    assert reporting.short_name_from_title(title) == expected


def test_safe_name():
    assert reporting.safe_name("Connection Setup") == "Connection_Setup"
    assert reporting.safe_name("a/b\\c:d") == "a_b_c_d"
    assert reporting.safe_name("***") == "run"


def test_every_test_module_has_a_panel_title():
    names = [reporting.panel_name(p) for p in sorted(TEST_DIR.glob("test_*.py"))]
    assert all(n and not n.startswith("test_") for n in names), names
    assert len(set(names)) == len(names), f"Duplicate panel names: {names}"


def test_history_appends_and_is_capped(tmp_path, monkeypatch):
    monkeypatch.setattr(reporting.RunReporter, "HISTORY_LIMIT", 3)
    rows = [{"result": "PASSED"}, {"result": "FAILED"}, {"result": "SKIPPED"}]
    for i in range(5):
        history = reporting.RunReporter._update_history(tmp_path, f"run{i}", rows)
    assert [h["run"] for h in history] == ["run2", "run3", "run4"]
    assert history[-1] == {"run": "run4", "total": 3, "passed": 1, "failed": 1, "skipped": 1}
    assert json.loads((tmp_path / "history.json").read_text()) == history


def test_write_creates_html_excel_and_summary(tmp_path, monkeypatch):
    shot = tmp_path / "Panel" / "fail_x.png"
    shot.parent.mkdir()
    shot.write_bytes(b"png")
    monkeypatch.setattr(reporting, "_environment", {"Tester status": "Connected"})
    rows = [
        {"node_id": "a::test_TC_01_ok", "panel": "Help", "title": "TC_01 · Ok", "spec": "HELP-01",
         "result": "PASSED", "duration": 1.0, "remarks": "", "message": "", "screenshot": "", "trace": ""},
        {"node_id": "a::test_TC_02_bad", "panel": "Help", "title": "TC_02 · Bad", "spec": "",
         "result": "FAILED", "duration": 2.0, "remarks": "boom", "message": "E  AssertionError: boom",
         "screenshot": str(shot), "trace": ""},
    ]
    reporting.RunReporter._write(tmp_path, "Help - now", rows, 3.0, [])

    page = (tmp_path / "report.html").read_text(encoding="utf-8")
    assert "TC_02 · Bad" in page and "AssertionError: boom" in page
    assert "src='Panel/fail_x.png'" in page, "screenshot must be linked relative to the report"
    assert "Tester status" in page

    wb = openpyxl.load_workbook(tmp_path / "report.xlsx")
    assert wb.sheetnames == ["Results", "Environment"]
    results = [[c.value for c in row] for row in wb["Results"].iter_rows(min_row=2)]
    assert [v or "" for v in results[1][2:5]] == ["TC_02 · Bad", "", "FAILED"]

    summary = json.loads((tmp_path / "summary.json").read_text(encoding="utf-8"))
    assert summary["counts"] == {"PASSED": 1, "FAILED": 1}
    assert summary["Passed"] == 1 and summary["Failed"] == 1
