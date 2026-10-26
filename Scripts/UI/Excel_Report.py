"""
Excel report for a GRL USB PD C2 automation run (openpyxl).
Sheets: Results (one row per test) and Environment (run / tester details).
"""
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
RESULT_FILLS = {
    "PASSED": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
    "FAILED": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "ERROR": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "SKIPPED": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    "XFAIL": PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid"),
}


def _header(ws, headers: list[str]) -> None:
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _autosize(ws, max_width: int = 80) -> None:
    for col in ws.columns:
        longest = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 3, 10), max_width)


def generate_excel_report(report_data: dict[str, Any], output_path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    _header(ws, ["S.No", "Panel", "Test", "Spec", "Result", "Duration", "Remarks", "Screenshot"])
    for row, tc in enumerate(report_data.get("Test Cases", []), 2):
        result = str(tc.get("TC_Result", "UNKNOWN")).upper()
        values = [row - 1, tc.get("Panel", ""), tc.get("Title") or tc.get("TC_ID", ""), tc.get("Spec", ""),
                  result, tc.get("TC_Execution_Duration", ""), tc.get("TC_Remarks", ""), tc.get("Screenshot", "")]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=col == 7)
        if result in RESULT_FILLS:
            ws.cell(row=row, column=5).fill = RESULT_FILLS[result]
        if tc.get("Screenshot"):
            ws.cell(row=row, column=8).hyperlink = tc["Screenshot"]
    _autosize(ws)

    ws_env = wb.create_sheet("Environment")
    _header(ws_env, ["Item", "Value"])
    summary = {
        "Report": report_data.get("Report Name", ""),
        "Total tests": report_data.get("Total Tests", 0),
        "Passed": report_data.get("Passed", 0),
        "Failed": report_data.get("Failed", 0),
        "Skipped / XFAIL": report_data.get("Skipped", 0),
        "Duration": report_data.get("Total Time", ""),
    }
    for row, (key, value) in enumerate({**summary, **report_data.get("Environment", {})}.items(), 2):
        ws_env.cell(row=row, column=1, value=key).font = Font(bold=True)
        ws_env.cell(row=row, column=2, value=str(value))
    _autosize(ws_env)

    wb.save(output_path)
